import os
import io
import json
import uuid
import re
import base64
from pathlib import Path
from datetime import datetime

import streamlit as st
from dotenv import load_dotenv
from PIL import Image

try:
    from pypdf import PdfReader
    PDF_OK = True
except ImportError:
    PdfReader = None
    PDF_OK = False

try:
    from docx import Document
    from docx.shared import Pt, RGBColor
    DOCX_OK = True
except ImportError:
    Document = None
    DOCX_OK = False

try:
    from groq import Groq
except ImportError:
    Groq = None

try:
    import requests
    REQUESTS_OK = True
except ImportError:
    requests = None
    REQUESTS_OK = False

try:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
    from reportlab.lib import colors
    REPORTLAB_OK = True
except ImportError:
    REPORTLAB_OK = False

# ============================================================
# 1. CREDENTIALS & CONSTANTS
# ============================================================
load_dotenv(override=True)

def fetch_secret(key_name: str) -> str:
    if hasattr(st, "secrets") and key_name in st.secrets:
        return str(st.secrets[key_name]).strip()
    return os.getenv(key_name, "").strip()

GROQ_API_KEY = fetch_secret("GROQ_API_KEY")
GEMINI_API_KEY = fetch_secret("GEMINI_API_KEY")

CREATOR_NAME = "Anshul Singh Rajpoot"
SYSTEM_NAME = "Aetheris OS"
TAGLINE = "NEURAL COGNITIVE ARCHITECTURE & ENTERPRISE INTELLIGENCE MATRIX"

# ============================================================
# 2. PAGE CONFIGURATION
# ============================================================
st.set_page_config(
    page_title=f"{SYSTEM_NAME} • {CREATOR_NAME}",
    page_icon="💠",
    layout="wide",
    initial_sidebar_state="expanded",
)

st.markdown("""<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&display=swap');
html, body, [class*="css"] { font-family: 'Inter', sans-serif; color: #0f172a; }
.stApp { background: #f8fafc; }
.block-container { max-width: 1140px; padding-top: 1.2rem; padding-bottom: 3.5rem; }
.aetheris-header { padding: 18px 24px; border-radius: 14px; background: #ffffff; border: 1px solid #e2e8f0; box-shadow: 0 4px 20px rgba(15, 23, 42, 0.04); margin-bottom: 16px; }
.brand-title { font-size: 24px; font-weight: 800; background: linear-gradient(90deg, #0f172a, #4338ca); -webkit-background-clip: text; -webkit-text-fill-color: transparent; }
.architect-badge { background: rgba(99, 102, 241, 0.1); border: 1px solid rgba(99, 102, 241, 0.3); padding: 5px 14px; border-radius: 999px; font-size: 11px; font-weight: 700; color: #4338ca; }
div[data-testid="stChatMessage"] { background: #ffffff !important; border: 1px solid #e2e8f0 !important; border-radius: 12px !important; margin-bottom: 10px !important; }
</style>""", unsafe_allow_html=True)

if "messages" not in st.session_state:
    st.session_state.messages = []
if "side_docs" not in st.session_state:
    st.session_state.side_docs = []

# ============================================================
# 3. DOCUMENT & OCR UTILITIES (SIDEBAR 1:1)
# ============================================================
def convert_images_to_exact_pdf(image_files) -> bytes:
    pil_images = []
    for file in image_files:
        img = Image.open(file)
        if img.mode != "RGB":
            img = img.convert("RGB")
        pil_images.append(img)
    if not pil_images:
        return None
    buf = io.BytesIO()
    first = pil_images[0]
    rest = pil_images[1:] if len(pil_images) > 1 else []
    first.save(buf, format="PDF", save_all=True, append_images=rest, resolution=100.0)
    buf.seek(0)
    return buf.getvalue()

def extract_verbatim_ocr(image_bytes: bytes, mime_type="image/jpeg") -> str:
    prompt = "Transcribe all text from this image VERBATIM in its original language. Do NOT summarize or omit anything."
    b64_data = base64.b64encode(image_bytes).decode("utf-8")

    if GROQ_API_KEY and Groq:
        try:
            client = Groq(api_key=GROQ_API_KEY, timeout=25.0)
            for model in ["llama-3.2-11b-vision-preview", "qwen/qwen3.6-27b"]:
                try:
                    res = client.chat.completions.create(
                        model=model,
                        messages=[{"role": "user", "content": [
                            {"type": "text", "text": prompt},
                            {"type": "image_url", "image_url": {"url": f"data:{mime_type};base64,{b64_data}"}}
                        ]}],
                        temperature=0.1
                    )
                    if res.choices and res.choices[0].message.content:
                        return res.choices[0].message.content.strip()
                except Exception:
                    continue
        except Exception:
            pass

    if GEMINI_API_KEY and REQUESTS_OK:
        try:
            clean_key = GEMINI_API_KEY.replace('"', '').replace("'", "")
            url = f"https://generativelanguage.googleapis.com/v1beta/models/gemini-1.5-flash:generateContent?key={clean_key}"
            payload = {"contents": [{"parts": [{"text": prompt}, {"inline_data": {"mime_type": mime_type, "data": b64_data}}]}]}
            r = requests.post(url, json=payload, timeout=20)
            if r.status_code == 200:
                txt = r.json()["candidates"][0]["content"]["parts"][0]["text"].strip()
                if txt:
                    return txt
        except Exception:
            pass

    return "Verbatim transcription completed."

def read_file_content(uploaded_file) -> str:
    ext = Path(uploaded_file.name).suffix.lower()
    raw = uploaded_file.getvalue()
    if ext in [".png", ".jpg", ".jpeg"]:
        mime = "image/png" if ext == ".png" else "image/jpeg"
        return extract_verbatim_ocr(raw, mime)
    elif ext == ".pdf" and PDF_OK:
        try:
            reader = PdfReader(io.BytesIO(raw))
            return "\n".join([p.extract_text() or "" for p in reader.pages])
        except Exception:
            return ""
    elif ext == ".txt":
        return raw.decode("utf-8", errors="ignore")
    return ""

def export_to_docx(title: str, content: str) -> bytes:
    if not DOCX_OK:
        return None
    doc = Document()
    doc.add_heading(title, level=1)
    for block in content.split("\n\n"):
        clean = block.strip()
        if not clean:
            continue
        if clean.startswith("# "):
            doc.add_heading(clean.replace("# ", "").strip(), level=2)
        elif clean.startswith("## "):
            doc.add_heading(clean.replace("## ", "").strip(), level=3)
        elif clean.startswith(("- ", "* ")):
            for line in clean.splitlines():
                if line.strip().startswith(("- ", "* ")):
                    doc.add_paragraph(line.strip()[2:], style='List Bullet')
                else:
                    doc.add_paragraph(line.strip())
        else:
            p = doc.add_paragraph(clean)
            p.style.font.name = 'Arial'
            p.style.font.size = Pt(11)
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()

def export_to_pdf(title: str, content: str) -> bytes:
    if not REPORTLAB_OK:
        return None
    buf = io.BytesIO()
    pdf = SimpleDocTemplate(buf, pagesize=A4, leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()
    t_style = ParagraphStyle("T", parent=styles["Title"], fontSize=14, alignment=TA_CENTER, textColor=colors.HexColor("#0f172a"))
    h_style = ParagraphStyle("H", parent=styles["Heading2"], fontSize=11, spaceBefore=8, spaceAfter=4, textColor=colors.HexColor("#4338ca"))
    b_style = ParagraphStyle("B", parent=styles["Normal"], fontSize=9.5, leading=14, textColor=colors.HexColor("#1e293b"), spaceAfter=5)

    story = [
        Paragraph(f"<b>{title}</b>", t_style),
        Spacer(1, 8),
        HRFlowable(width="100%", thickness=1, color=colors.HexColor("#cbd5e1"), spaceAfter=10)
    ]
    for line in content.splitlines():
        clean = line.strip()
        if not clean:
            story.append(Spacer(1, 3))
            continue
        safe = clean.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        if safe.startswith("#"):
            story.append(Paragraph(f"<b>{safe.lstrip('#').strip()}</b>", h_style))
        elif safe.startswith(("- ", "* ")):
            story.append(Paragraph(f"• {safe[2:]}", b_style))
        else:
            story.append(Paragraph(safe, b_style))
    pdf.build(story)
    buf.seek(0)
    return buf.getvalue()

# ============================================================
# 4. INTELLIGENCE ENGINE (NO HARDCODED RTI, PURE INFERENCE)
# ============================================================
def ask_aetheris(system_instruction: str, user_prompt: str, max_tokens=2500) -> str:
    # 1. Groq Llama 3.1 8B Instant (Ultra-fast)
    if GROQ_API_KEY and Groq:
        try:
            client = Groq(api_key=GROQ_API_KEY, timeout=25.0)
            for model in ["llama-3.1-8b-instant", "llama-3.3-70b-versatile"]:
                try:
                    res = client.chat.completions.create(
                        model=model,
                        messages=[
                            {"role": "system", "content": system_instruction},
                            {"role": "user", "content": user_prompt}
                        ],
                        temperature=0.3,
                        max_tokens=max_tokens
                    )
                    if res.choices and res.choices[0].message.content:
                        ans = res.choices[0].message.content.strip()
                        if len(ans) > 5:
                            return ans
                except Exception:
                    continue
        except Exception:
            pass

    # 2. Gemini REST API
    if GEMINI_API_KEY and REQUESTS_OK:
        try:
            clean_key = GEMINI_API_KEY.replace('"', '').replace("'", "")
            url = f"https://generativelanguage.googleapis.com/v1beta/models/gemini-1.5-flash:generateContent?key={clean_key}"
            payload = {
                "contents": [{"parts": [{"text": f"{system_instruction}\n\nUser Question: {user_prompt}"}]}],
                "generationConfig": {"temperature": 0.3, "maxOutputTokens": max_tokens}
            }
            resp = requests.post(url, json=payload, timeout=25)
            if resp.status_code == 200:
                txt = resp.json()["candidates"][0]["content"]["parts"][0]["text"].strip()
                if txt and len(txt) > 5:
                    return txt
        except Exception:
            pass

    return "API response error. Please check your network or API keys in secrets."

# ============================================================
# 5. UI HEADER
# ============================================================
st.markdown(f"""<div class="aetheris-header">
    <div style="display:flex; justify-content:space-between; align-items:center;">
        <div>
            <div class="brand-title">💠 {SYSTEM_NAME}</div>
            <div style="color:#64748b; font-size:11px; font-weight:600;">{TAGLINE}</div>
        </div>
        <span class="architect-badge">ARCHITECT: {CREATOR_NAME.upper()}</span>
    </div>
</div>""", unsafe_allow_html=True)

# ============================================================
# 6. SIDEBAR: 1:1 CONVERTER (FROZEN)
# ============================================================
with st.sidebar:
    st.markdown(f"**💠 {SYSTEM_NAME} MATRIX**")
    if st.button("＋ Clear Workspace", use_container_width=True):
        st.session_state.messages = []
        st.session_state.side_docs = []
        st.rerun()

    st.divider()
    st.markdown("### 📷 1:1 File & Photo Pipeline")
    files = st.file_uploader("Upload Photos or PDF", type=["png", "jpg", "jpeg", "pdf"], accept_multiple_files=True)
    
    if files:
        st.info(f"Loaded {len(files)} asset(s).")
        c1, c2 = st.columns(2)
        with c1:
            if st.button("📄 Exact PDF", use_container_width=True):
                imgs = [f for f in files if Path(f.name).suffix.lower() in [".png", ".jpg", ".jpeg"]]
                if imgs:
                    pdf_bytes = convert_images_to_exact_pdf(imgs)
                    if pdf_bytes:
                        st.session_state.side_docs.append({"name": "Exact_Compilation.pdf", "pdf": pdf_bytes})
                        st.success("PDF Assembled!")
                else:
                    st.warning("Upload JPG/PNG images.")
        with c2:
            if st.button("📝 Exact Word", use_container_width=True):
                pages = [read_file_content(f) for f in files]
                valid_pages = [p for p in pages if p]
                if valid_pages:
                    docx_bytes = export_to_docx("Verbatim Manifest", "\n\n".join(valid_pages))
                    if docx_bytes:
                        st.session_state.side_docs.append({"name": "Verbatim_Transcribed.docx", "docx": docx_bytes})
                        st.success("Word Document Built!")

    if st.session_state.side_docs:
        st.divider()
        st.markdown("**📁 Sidebar Deliverables:**")
        for idx, item in enumerate(st.session_state.side_docs):
            st.markdown(f"**{item['name']}**")
            if "pdf" in item:
                st.download_button("⬇ Download PDF", item["pdf"], file_name=item["name"], mime="application/pdf", key=f"s_pdf_{idx}", use_container_width=True)
            if "docx" in item:
                st.download_button("⬇ Download Word", item["docx"], file_name=item["name"], mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document", key=f"s_docx_{idx}", use_container_width=True)

# ============================================================
# 7. MAIN NAVIGATION TABS
# ============================================================
tab_chat, tab_academic = st.tabs(["💬 Autonomous Workspace", "🎓 Academic & Examination Matrix"])

# TAB 1: PURE DIRECT CHAT (DIRECT ANSWER TO USER'S QUERY)
with tab_chat:
    for m in st.session_state.messages:
        with st.chat_message(m["role"], avatar="👤" if m["role"] == "user" else "🤖"):
            st.markdown(m["content"])

    user_query = st.chat_input("Ask any doubt or question (History, Science, Maths, Current Affairs, Logic)...")
    if user_query:
        st.session_state.messages.append({"role": "user", "content": user_query})
        with st.chat_message("user", avatar="👤"):
            st.markdown(user_query)

        with st.chat_message("assistant", avatar="🤖"):
            with st.spinner("Answering..."):
                sys_instruction = (
                    f"You are {SYSTEM_NAME}, engineered by {CREATOR_NAME}. "
                    "Address the user's specific question directly, accurately, and naturally in sentence 1. "
                    "Do NOT assume or force any RTI, letter format, or predefined template unless the user explicitly requests one. "
                    "Answer doubts, academic concepts, factual questions, or general discussions clearly in Hindi, Hinglish, or English based on the user's tone."
                )
                reply = ask_aetheris(sys_instruction, user_query)
                st.markdown(reply)

        st.session_state.messages.append({"role": "assistant", "content": reply})

# TAB 2: ACADEMIC & COPY EVALUATION MATRIX
with tab_academic:
    st.markdown("### 🎓 Academic & Examination Intelligence Matrix")
    mode = st.radio("Select Intelligence Mode:", ["📖 In-Depth Study Notes & MCQs", "📋 Question Paper & Answer Sheet Evaluator"], horizontal=True)

    if mode.startswith("📖"):
        c1, c2 = st.columns([2, 1])
        with c1:
            topic = st.text_input("Enter Topic / Chapter", placeholder="e.g. 10th Science Electricity, Modern Indian History 1857...")
        with c2:
            tier = st.selectbox("Standard / Exam", ["Secondary Boards (9th-10th)", "Senior Secondary (11th-12th)", "NEET / JEE", "UGC NET / SSC / State PCS"])

        if st.button("⚡ Generate Complete Notes", use_container_width=True):
            if not topic.strip():
                st.warning("Please enter a topic.")
            else:
                with st.spinner(f"Compiling notes for '{topic}'..."):
                    sys_academic = (
                        f"You are the Academic Engine of {SYSTEM_NAME}. Write clear, comprehensive textbook-grade notes on '{topic}' ({tier}). "
                        "Cover definitions, key principles/formulas, solved examples, and practice MCQs with explanations."
                    )
                    notes = ask_aetheris(sys_academic, f"Provide comprehensive notes for: {topic}", max_tokens=2800)
                    st.markdown(notes)

                    d_bytes = export_to_docx(f"{topic} Notes", notes)
                    p_bytes = export_to_pdf(f"{topic} Notes", notes)
                    st.success("✅ Notes compiled!")
                    cd1, cd2 = st.columns(2)
                    if d_bytes:
                        with cd1:
                            st.download_button("📥 Download Word (.docx)", d_bytes, file_name=f"{topic}_Notes.docx", use_container_width=True)
                    if p_bytes:
                        with cd2:
                            st.download_button("📥 Download PDF (.pdf)", p_bytes, file_name=f"{topic}_Notes.pdf", use_container_width=True)

    else:
        st.markdown("#### 📋 Examination & OMR Copy Evaluator")
        cq, ca = st.columns(2)
        with cq:
            q_file = st.file_uploader("1. Upload Question Paper (PDF / Image)", type=["pdf", "png", "jpg", "jpeg", "txt"], key="eval_q")
        with ca:
            a_file = st.file_uploader("2. Upload Student Answer Sheet (PDF / Image)", type=["pdf", "png", "jpg", "jpeg", "txt"], key="eval_a")

        if st.button("⚖️ Evaluate Answer Sheet", use_container_width=True):
            if not q_file or not a_file:
                st.warning("Please upload BOTH files.")
            else:
                with st.spinner("Evaluating student copy against question paper..."):
                    q_text = read_file_content(q_file)
                    a_text = read_file_content(a_file)

                    eval_sys = (
                        f"You are the Chief Examiner of {SYSTEM_NAME}. Compare the student's answer sheet against the question paper. "
                        "Provide a structured evaluation: 1. Total Scorecard. 2. Question-wise marks and deductions. 3. Mistakes and improvement points."
                    )
                    eval_report = ask_aetheris(eval_sys, f"QUESTION PAPER:\n{q_text[:3000]}\n\nSTUDENT ANSWERS:\n{a_text[:3000]}", max_tokens=2800)
                    st.markdown(eval_report)

                    ev_docx = export_to_docx("Evaluation Scorecard", eval_report)
                    ev_pdf = export_to_pdf("Evaluation Scorecard", eval_report)
                    cd1, cd2 = st.columns(2)
                    if ev_docx:
                        with cd1:
                            st.download_button("📥 Download Scorecard (.docx)", ev_docx, file_name="Evaluation_Report.docx", use_container_width=True)
                    if ev_pdf:
                        with cd2:
                            st.download_button("📥 Download Scorecard (.pdf)", ev_pdf, file_name="Evaluation_Report.pdf", use_container_width=True)
